"""Generate a polished WMS Project Presentation as a .pptx file.

Saves the file to the user's Downloads folder.
Run with the project virtualenv:
    .venv\Scripts\python.exe generate_pptx.py
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
DARK_BG = RGBColor(0x0F, 0x17, 0x2A)        # slate-900
DARKER_BG = RGBColor(0x1E, 0x29, 0x3B)      # slate-800
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
BRAND = RGBColor(0x25, 0x63, 0xEB)          # blue-600
BRAND_LIGHT = RGBColor(0x93, 0xC5, 0xFD)    # blue-300
ACCENT = RGBColor(0x06, 0xB6, 0xD4)         # cyan-500
PURPLE = RGBColor(0x7C, 0x3A, 0xED)         # violet-600
GOOD = RGBColor(0x10, 0xB9, 0x81)           # emerald-500
WARN = RGBColor(0xF5, 0x9E, 0x0B)           # amber-500
DANGER = RGBColor(0xEF, 0x44, 0x44)         # red-500
GRAY = RGBColor(0x6B, 0x72, 0x80)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_slide_bg(slide, color: RGBColor) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font_size: int = 14,
    color: RGBColor = WHITE,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font_name: str = "Segoe UI",
):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor

    if isinstance(text, str):
        lines = text.split("\n")
    else:
        lines = list(text)

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb


def add_bullets(slide, left, top, width, height, bullets, *, font_size=14, color=LIGHT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.name = "Segoe UI"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return tb


def add_card(
    slide,
    left,
    top,
    width,
    height,
    title: str,
    bullets: list[str],
    *,
    accent: RGBColor = BRAND,
    title_color: RGBColor = WHITE,
):
    add_rect(slide, left, top, width, height, DARKER_BG, line_color=accent)
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()
    accent_bar.shadow.inherit = False

    add_text_box(
        slide,
        left + Inches(0.2),
        top + Inches(0.1),
        width - Inches(0.3),
        Inches(0.4),
        title,
        font_size=16,
        bold=True,
        color=title_color,
    )
    add_bullets(
        slide,
        left + Inches(0.2),
        top + Inches(0.55),
        width - Inches(0.3),
        height - Inches(0.6),
        bullets,
        font_size=12,
    )


def add_slide_title(slide, title: str, subtitle: str | None = None):
    add_text_box(
        slide,
        Inches(0.5),
        Inches(0.3),
        Inches(12.3),
        Inches(0.6),
        title,
        font_size=32,
        bold=True,
        color=WHITE,
    )
    if subtitle:
        add_text_box(
            slide,
            Inches(0.5),
            Inches(0.95),
            Inches(12.3),
            Inches(0.4),
            subtitle,
            font_size=16,
            italic=True,
            color=BRAND_LIGHT,
        )
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.45), Inches(2), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    line.shadow.inherit = False


def add_slide_number(slide, current: int, total: int):
    add_text_box(
        slide,
        Inches(12.3),
        Inches(7.05),
        Inches(1),
        Inches(0.3),
        f"{current} / {total}",
        font_size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_footer(slide):
    add_text_box(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(10),
        Inches(0.3),
        "WMS — Warehouse Management System  |  FastAPI · React · TypeScript · SQLite",
        font_size=10,
        color=MUTED,
    )


# ---------------------------------------------------------------------------
# Build presentation
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
TOTAL_SLIDES = 14


# ============ SLIDE 1 — TITLE ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)

# Decorative accent bars
bar1 = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15)
)
bar1.fill.solid(); bar1.fill.fore_color.rgb = BRAND; bar1.line.fill.background()
bar1.shadow.inherit = False

bar2 = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.35), Inches(13.333), Inches(0.15)
)
bar2.fill.solid(); bar2.fill.fore_color.rgb = PURPLE; bar2.line.fill.background()
bar2.shadow.inherit = False

# Title
add_text_box(
    slide, Inches(0.5), Inches(2.0), Inches(12.3), Inches(1.2),
    "Warehouse Management System",
    font_size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
)
add_text_box(
    slide, Inches(0.5), Inches(3.1), Inches(12.3), Inches(0.6),
    "Full-Stack Order & AGV Automation Platform",
    font_size=22, italic=True, color=BRAND_LIGHT, align=PP_ALIGN.CENTER,
)

# Three pillars
pillars = [
    ("Warehouse", "Smart Storage", BRAND),
    ("Orders", "Auto-Routed", PURPLE),
    ("AGVs", "Automated", ACCENT),
]
card_w = Inches(2.4); card_h = Inches(1.4); gap = Inches(0.4)
total_w = card_w * 3 + gap * 2
start_left = (Inches(13.333) - total_w) / 2
for i, (title, sub, color) in enumerate(pillars):
    left = start_left + (card_w + gap) * i
    add_rect(slide, left, Inches(4.4), card_w, card_h, DARKER_BG, line_color=color)
    add_text_box(slide, left, Inches(4.6), card_w, Inches(0.5), title,
                 font_size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, left, Inches(5.15), card_w, Inches(0.4), sub,
                 font_size=14, color=MUTED, align=PP_ALIGN.CENTER)

# Presenter info
add_text_box(
    slide, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.4),
    "Presented by: ____________________     Date: ____________________",
    font_size=14, color=MUTED, align=PP_ALIGN.CENTER,
)


# ============ SLIDE 2 — AGENDA ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Agenda", "What we'll cover today")

agenda = [
    ("1", "The Problem", "Why warehouses need automation"),
    ("2", "The Solution", "Our full-stack WMS approach"),
    ("3", "Architecture", "How the pieces fit together"),
    ("4", "Tech Stack", "Tools and technologies used"),
    ("5", "Core Features", "What the system does"),
    ("6", "Live Demo", "See it in action"),
    ("7", "Challenges", "What we learned"),
    ("8", "What's Next", "Future roadmap"),
]
col_w = Inches(6.0); row_h = Inches(1.0); gap_x = Inches(0.4); gap_y = Inches(0.15)
start_top = Inches(1.9)
for i, (num, title, desc) in enumerate(agenda):
    col = i % 2
    row = i // 2
    left = Inches(0.5) + col * (col_w + gap_x)
    top = start_top + row * (row_h + gap_y)
    add_rect(slide, left, top, col_w, row_h, DARKER_BG, line_color=BRAND)
    add_text_box(slide, left + Inches(0.2), top + Inches(0.15), Inches(0.6), Inches(0.7),
                 num, font_size=28, bold=True, color=ACCENT)
    add_text_box(slide, left + Inches(0.85), top + Inches(0.15), col_w - Inches(1),
                 Inches(0.4), title, font_size=18, bold=True, color=WHITE)
    add_text_box(slide, left + Inches(0.85), top + Inches(0.55), col_w - Inches(1),
                 Inches(0.4), desc, font_size=12, color=MUTED)
add_slide_number(slide, 2, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 3 — THE PROBLEM ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "The Problem", "Why warehouses need a WMS")

add_card(
    slide, Inches(0.5), Inches(2.0), Inches(6.0), Inches(2.5),
    "Manual chaos",
    [
        "Hundreds of orders, thousands of items",
        "Wrong items shipped, lost stock",
        "Robots idle while orders pile up",
        "Human errors and delays",
    ],
    accent=WARN,
)

add_card(
    slide, Inches(6.85), Inches(2.0), Inches(6.0), Inches(2.5),
    "No visibility",
    [
        "Managers fly blind",
        "No real-time order status",
        "No audit of who did what, when",
        "Reports take days, not seconds",
    ],
    accent=DANGER,
)

# Quote
quote_box = add_rect(slide, Inches(1.5), Inches(5.0), Inches(10.3), Inches(1.2),
                     DARKER_BG, line_color=ACCENT)
add_text_box(
    slide, Inches(1.8), Inches(5.25), Inches(9.8), Inches(0.7),
    '"Without coordination, even the best warehouse becomes the bottleneck."',
    font_size=18, italic=True, color=LIGHT, align=PP_ALIGN.CENTER,
    anchor=MSO_ANCHOR.MIDDLE,
)

add_slide_number(slide, 3, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 4 — THE SOLUTION ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "The Solution", "One system. Four pillars.")

solutions = [
    ("Manage Everything", "Products, Orders, Zones, AGVs — full CRUD in one place.", BRAND),
    ("Secure Access", "JWT-protected login on every endpoint.", PURPLE),
    ("Auto-Assignment", "Idle AGV picks the next order automatically.", ACCENT),
    ("Real-Time Tracking", "Live status updates + complete audit log.", GOOD),
]
card_w = Inches(6.0); card_h = Inches(2.3); gap_x = Inches(0.35); gap_y = Inches(0.3)
start_top = Inches(1.9)
for i, (title, desc, color) in enumerate(solutions):
    col = i % 2; row = i // 2
    left = Inches(0.5) + col * (card_w + gap_x)
    top = start_top + row * (card_h + gap_y)
    add_rect(slide, left, top, card_w, card_h, DARKER_BG, line_color=color)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.1), card_h)
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    bar.shadow.inherit = False
    add_text_box(slide, left + Inches(0.3), top + Inches(0.3), card_w - Inches(0.4),
                 Inches(0.6), title, font_size=20, bold=True, color=WHITE)
    add_text_box(slide, left + Inches(0.3), top + Inches(1.0), card_w - Inches(0.4),
                 Inches(1.2), desc, font_size=14, color=LIGHT)
add_slide_number(slide, 4, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 5 — ARCHITECTURE ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Architecture", "Three layers + one simulator")

# Three boxes in a row
arch_top = Inches(2.3); arch_h = Inches(1.5); arch_w = Inches(2.8)
gap_arrow = Inches(0.6)
total_arch_w = arch_w * 3 + gap_arrow * 2
start_left = (Inches(13.333) - total_arch_w) / 2

arch_items = [
    ("React Web", "Frontend", ACCENT),
    ("FastAPI", "Backend", BRAND),
    ("SQLite", "Database", PURPLE),
]
positions = []
for i, (name, label, color) in enumerate(arch_items):
    left = start_left + i * (arch_w + gap_arrow)
    positions.append(left)
    add_rect(slide, left, arch_top, arch_w, arch_h, color)
    add_text_box(slide, left, arch_top + Inches(0.3), arch_w, Inches(0.5),
                 name, font_size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, left, arch_top + Inches(0.85), arch_w, Inches(0.4),
                 label, font_size=14, color=WHITE, align=PP_ALIGN.CENTER)

# Arrows between boxes
for i in range(2):
    arrow_left = positions[i] + arch_w
    add_text_box(
        slide, arrow_left, arch_top + Inches(0.45), gap_arrow, Inches(0.6),
        "<-->", font_size=24, bold=True, color=MUTED, align=PP_ALIGN.CENTER,
    )

# Simulator box
sim_w = Inches(5.5); sim_h = Inches(0.9)
sim_left = (Inches(13.333) - sim_w) / 2
sim_top = arch_top + arch_h + Inches(0.6)
add_rect(slide, sim_left, sim_top, sim_w, sim_h, WARN)
add_text_box(slide, sim_left, sim_top + Inches(0.2), sim_w, Inches(0.5),
             "Background AGV Simulator (async lifespan task)",
             font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Explanations
desc_top = sim_top + sim_h + Inches(0.4)
descs = [
    ("Frontend", "What the user sees", ACCENT),
    ("Backend", "The brain — APIs & rules", BRAND),
    ("Database", "Persistent memory", PURPLE),
]
for i, (t, d, c) in enumerate(descs):
    left = positions[i]
    add_text_box(slide, left, desc_top, arch_w, Inches(0.4),
                 t, font_size=14, bold=True, color=c, align=PP_ALIGN.CENTER)
    add_text_box(slide, left, desc_top + Inches(0.4), arch_w, Inches(0.4),
                 d, font_size=12, color=MUTED, align=PP_ALIGN.CENTER)

add_slide_number(slide, 5, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 6 — TECH STACK ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Tech Stack", "Modern, type-safe, production-ready")

add_card(
    slide, Inches(0.5), Inches(2.0), Inches(6.0), Inches(4.0),
    "Backend",
    [
        "FastAPI — modern async Python framework",
        "SQLAlchemy — ORM for clean DB access",
        "SQLite — embedded persistent storage",
        "JWT — secure token-based authentication",
        "Pydantic — automatic input validation",
        "Uvicorn — lightning-fast ASGI server",
    ],
    accent=BRAND,
)

add_card(
    slide, Inches(6.85), Inches(2.0), Inches(6.0), Inches(4.0),
    "Frontend",
    [
        "React 18 — component-based UI",
        "TypeScript — type safety end-to-end",
        "Vite — fast dev server & bundler",
        "Axios — typed HTTP client",
        "Context API — auth state management",
        "React Router — protected routes",
    ],
    accent=ACCENT,
)

add_text_box(
    slide, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.5),
    "Why this stack? — Fast to build, type-safe end-to-end, production-ready.",
    font_size=14, italic=True, color=BRAND_LIGHT, align=PP_ALIGN.CENTER,
)
add_slide_number(slide, 6, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 7 — CORE ENTITIES ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Core Entities", "Six entities model the entire warehouse")

entities = [
    ("Product", "Items stored in warehouse", "Full CRUD"),
    ("Order", "Customer requests + priority", "Create / Cancel / Confirm"),
    ("Zone", "Source & destination locations", "Full CRUD"),
    ("AGV", "Automated guided vehicles", "Status & assignment"),
    ("User", "Authenticated staff", "Login / Auth"),
    ("SystemLog", "Complete audit trail", "Read-only history"),
]

# Header row
header_top = Inches(2.0); row_h = Inches(0.55)
table_left = Inches(0.7); table_w = Inches(11.9)
col1 = Inches(2.5); col2 = Inches(5.5); col3 = Inches(3.9)

add_rect(slide, table_left, header_top, table_w, row_h, BRAND)
add_text_box(slide, table_left + Inches(0.2), header_top + Inches(0.1), col1, Inches(0.4),
             "Entity", font_size=14, bold=True, color=WHITE)
add_text_box(slide, table_left + col1 + Inches(0.2), header_top + Inches(0.1), col2,
             Inches(0.4), "Purpose", font_size=14, bold=True, color=WHITE)
add_text_box(slide, table_left + col1 + col2 + Inches(0.2), header_top + Inches(0.1),
             col3, Inches(0.4), "Operations", font_size=14, bold=True, color=WHITE)

for i, (name, purpose, ops) in enumerate(entities):
    top = header_top + row_h * (i + 1)
    bg = DARKER_BG if i % 2 == 0 else DARK_BG
    add_rect(slide, table_left, top, table_w, row_h, bg)
    add_text_box(slide, table_left + Inches(0.2), top + Inches(0.13), col1, Inches(0.4),
                 name, font_size=13, bold=True, color=BRAND_LIGHT)
    add_text_box(slide, table_left + col1 + Inches(0.2), top + Inches(0.13), col2,
                 Inches(0.4), purpose, font_size=12, color=LIGHT)
    add_text_box(slide, table_left + col1 + col2 + Inches(0.2), top + Inches(0.13),
                 col3, Inches(0.4), ops, font_size=12, color=MUTED)

add_slide_number(slide, 7, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 8 — ORDER LIFECYCLE ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Order Lifecycle", "From click → delivery, fully automated")

# Status pills
statuses = [
    ("pending", GRAY),
    ("validated", WARN),
    ("assigned", BRAND),
    ("in-transit", PURPLE),
    ("delivered", GOOD),
]
pill_w = Inches(1.7); pill_h = Inches(0.55); arrow_w = Inches(0.4)
total_w = pill_w * 5 + arrow_w * 4
start_left = (Inches(13.333) - total_w) / 2
pill_top = Inches(2.3)

for i, (status, color) in enumerate(statuses):
    left = start_left + i * (pill_w + arrow_w)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  left, pill_top, pill_w, pill_h)
    pill.fill.solid(); pill.fill.fore_color.rgb = color
    pill.line.fill.background()
    pill.shadow.inherit = False
    pill.adjustments[0] = 0.5
    add_text_box(slide, left, pill_top + Inches(0.1), pill_w, Inches(0.4),
                 status, font_size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER)
    if i < 4:
        add_text_box(slide, left + pill_w, pill_top + Inches(0.1), arrow_w, Inches(0.4),
                     ">", font_size=22, bold=True, color=MUTED,
                     align=PP_ALIGN.CENTER)

add_card(
    slide, Inches(0.5), Inches(3.5), Inches(6.0), Inches(2.8),
    "User actions",
    [
        "Create order with priority",
        "Pick destination zone",
        "Submit",
    ],
    accent=ACCENT,
)
add_card(
    slide, Inches(6.85), Inches(3.5), Inches(6.0), Inches(2.8),
    "System actions (automatic)",
    [
        "Validate & persist to database",
        "Find idle AGV",
        "Pick → Deliver → Log every step",
    ],
    accent=GOOD,
)
add_slide_number(slide, 8, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 9 — HOW API WORKS ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "How the API Connects Frontend & Backend",
                "One request — many layers")

# Steps
steps = [
    ("1", "User clicks 'Add Product'", "React captures form data", ACCENT),
    ("2", "api.post('/products/', data)", "Axios sends HTTP POST + JWT", BRAND),
    ("3", "FastAPI route receives", "Validates JWT + Pydantic schema", PURPLE),
    ("4", "Controller runs logic", "Builds object, saves to DB", WARN),
    ("5", "SQLAlchemy → SQLite", "INSERT SQL executes", GOOD),
    ("6", "JSON response returns", "React updates the UI", ACCENT),
]

step_top = Inches(2.0); step_h = Inches(0.7); gap_y = Inches(0.1)
for i, (num, title, desc, color) in enumerate(steps):
    top = step_top + i * (step_h + gap_y)
    add_rect(slide, Inches(0.5), top, Inches(12.3), step_h, DARKER_BG, line_color=color)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), top,
                                 Inches(0.1), step_h)
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False

    add_text_box(slide, Inches(0.7), top + Inches(0.1), Inches(0.6), Inches(0.5),
                 num, font_size=24, bold=True, color=color)
    add_text_box(slide, Inches(1.4), top + Inches(0.05), Inches(5.5), Inches(0.4),
                 title, font_size=14, bold=True, color=WHITE)
    add_text_box(slide, Inches(1.4), top + Inches(0.4), Inches(5.5), Inches(0.3),
                 desc, font_size=11, color=MUTED)
    add_text_box(slide, Inches(7.0), top + Inches(0.2), Inches(5.5), Inches(0.4),
                 "Frontend" if i < 2 else ("Backend" if i < 5 else "Frontend"),
                 font_size=12, italic=True, color=BRAND_LIGHT, align=PP_ALIGN.RIGHT)

add_slide_number(slide, 9, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 10 — LIVE DEMO ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Live Demo Time", "Switch to the running app")

# Big LIVE banner
banner_left = Inches(1.5); banner_top = Inches(2.0)
banner_w = Inches(10.3); banner_h = Inches(4.5)
add_rect(slide, banner_left, banner_top, banner_w, banner_h, DARKER_BG, line_color=WARN)

# Live indicator
dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.0), Inches(2.4),
                             Inches(0.3), Inches(0.3))
dot.fill.solid(); dot.fill.fore_color.rgb = DANGER
dot.line.fill.background()
dot.shadow.inherit = False
add_text_box(slide, Inches(2.4), Inches(2.3), Inches(2), Inches(0.5),
             "LIVE", font_size=24, bold=True, color=DANGER)

demo_steps = [
    "1.   Login → Dashboard",
    "2.   Create & edit a Product",
    "3.   Show Zones (source vs destination)",
    "4.   Create an Order with priority",
    "5.   Watch AGV auto-pick & deliver",
    "6.   Open Logs — full audit trail",
    "7.   Bonus: Swagger UI at /docs",
]
add_bullets(slide, Inches(2.0), Inches(3.0), Inches(9), Inches(3.4),
            demo_steps, font_size=18)

add_text_box(slide, Inches(0.5), Inches(6.7), Inches(12.3), Inches(0.4),
             "Time: 5–7 minutes  |  Switch to browser now",
             font_size=14, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

add_slide_number(slide, 10, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 11 — BY THE NUMBERS ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "By the Numbers", "What was built")

stats = [
    ("6", "Core Entities", BRAND),
    ("25+", "REST Endpoints", ACCENT),
    ("7", "React Pages", PURPLE),
    ("100%", "Type-Safe", GOOD),
    ("JWT", "Auth Secured", WARN),
    ("∞", "Audit Logged", BRAND_LIGHT),
]
card_w = Inches(3.9); card_h = Inches(2.0); gap_x = Inches(0.3); gap_y = Inches(0.3)
start_top = Inches(2.0)
for i, (num, label, color) in enumerate(stats):
    col = i % 3; row = i // 3
    left = Inches(0.55) + col * (card_w + gap_x)
    top = start_top + row * (card_h + gap_y)
    add_rect(slide, left, top, card_w, card_h, DARKER_BG, line_color=color)
    add_text_box(slide, left, top + Inches(0.4), card_w, Inches(0.9),
                 num, font_size=48, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text_box(slide, left, top + Inches(1.4), card_w, Inches(0.5),
                 label, font_size=14, color=MUTED, align=PP_ALIGN.CENTER)

add_slide_number(slide, 11, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 12 — CHALLENGES ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "Challenges & Learnings", "What this project taught me")

challenges = [
    ("Async background simulator",
     "Running a continuous task without blocking the API. Solved with FastAPI's lifespan handler.",
     WARN),
    ("Entity relationships",
     "Modeling Orders ↔ AGVs ↔ Zones cleanly so each can be queried independently.",
     BRAND),
    ("End-to-end auth",
     "JWT generation on backend, secure storage and refresh on the React side via Context.",
     PURPLE),
    ("Layered error handling",
     "Validation at schema layer, business rules at service, global handler for the unexpected.",
     ACCENT),
]
card_w = Inches(6.0); card_h = Inches(2.2); gap_x = Inches(0.35); gap_y = Inches(0.3)
start_top = Inches(2.0)
for i, (title, desc, color) in enumerate(challenges):
    col = i % 2; row = i // 2
    left = Inches(0.5) + col * (card_w + gap_x)
    top = start_top + row * (card_h + gap_y)
    add_rect(slide, left, top, card_w, card_h, DARKER_BG, line_color=color)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.1), card_h)
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    add_text_box(slide, left + Inches(0.3), top + Inches(0.25), card_w - Inches(0.4),
                 Inches(0.5), title, font_size=16, bold=True, color=WHITE)
    add_text_box(slide, left + Inches(0.3), top + Inches(0.85), card_w - Inches(0.4),
                 Inches(1.3), desc, font_size=12, color=LIGHT)

add_slide_number(slide, 12, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 13 — WHAT'S NEXT ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)
add_slide_title(slide, "What's Next", "Production-ready foundation")

next_items = [
    ("Scale", "Swap SQLite → PostgreSQL with one config change.", BRAND),
    ("Real Robots", "Replace simulator with actual AGV vendor APIs.", ACCENT),
    ("Reporting", "Add analytics, charts, mobile-friendly views.", PURPLE),
    ("Deploy", "Containerize with Docker, deploy to cloud.", GOOD),
    ("Notifications", "WebSocket-based live updates & alerts.", WARN),
    ("Tests", "Add unit + integration coverage with pytest.", DANGER),
]
card_w = Inches(3.9); card_h = Inches(2.0); gap_x = Inches(0.3); gap_y = Inches(0.3)
start_top = Inches(2.0)
for i, (title, desc, color) in enumerate(next_items):
    col = i % 3; row = i // 3
    left = Inches(0.55) + col * (card_w + gap_x)
    top = start_top + row * (card_h + gap_y)
    add_rect(slide, left, top, card_w, card_h, DARKER_BG, line_color=color)
    add_text_box(slide, left + Inches(0.2), top + Inches(0.2), card_w - Inches(0.4),
                 Inches(0.5), title, font_size=18, bold=True, color=color)
    add_text_box(slide, left + Inches(0.2), top + Inches(0.8), card_w - Inches(0.4),
                 Inches(1.1), desc, font_size=12, color=LIGHT)

add_slide_number(slide, 13, TOTAL_SLIDES); add_footer(slide)


# ============ SLIDE 14 — THANK YOU ============
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, DARK_BG)

bar1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                              Inches(13.333), Inches(0.15))
bar1.fill.solid(); bar1.fill.fore_color.rgb = BRAND; bar1.line.fill.background()
bar1.shadow.inherit = False
bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.35),
                              Inches(13.333), Inches(0.15))
bar2.fill.solid(); bar2.fill.fore_color.rgb = PURPLE; bar2.line.fill.background()
bar2.shadow.inherit = False

add_text_box(slide, Inches(0.5), Inches(1.5), Inches(12.3), Inches(1.2),
             "Thank You", font_size=72, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0.5), Inches(2.8), Inches(12.3), Inches(0.6),
             "Questions?", font_size=24, italic=True, color=BRAND_LIGHT,
             align=PP_ALIGN.CENTER)

# Quote
quote_left = Inches(1.5); quote_top = Inches(3.8)
quote_w = Inches(10.3); quote_h = Inches(1.3)
add_rect(slide, quote_left, quote_top, quote_w, quote_h, DARKER_BG, line_color=ACCENT)
add_text_box(slide, quote_left + Inches(0.3), quote_top + Inches(0.2),
             quote_w - Inches(0.6), Inches(0.45),
             '"From a user creating an order to a robot delivering it',
             font_size=18, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)
add_text_box(slide, quote_left + Inches(0.3), quote_top + Inches(0.7),
             quote_w - Inches(0.6), Inches(0.45),
             'the entire loop is automated."',
             font_size=18, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# URLs
url_top = Inches(5.5)
urls = [
    ("Backend", "localhost:8000"),
    ("API Docs", "localhost:8000/docs"),
    ("Frontend", "localhost:5173"),
]
url_w = Inches(3.5); gap = Inches(0.4)
total_url_w = url_w * 3 + gap * 2
start_url = (Inches(13.333) - total_url_w) / 2
for i, (label, url) in enumerate(urls):
    left = start_url + i * (url_w + gap)
    add_rect(slide, left, url_top, url_w, Inches(1.0), DARKER_BG, line_color=BRAND)
    add_text_box(slide, left, url_top + Inches(0.15), url_w, Inches(0.35),
                 label, font_size=12, color=MUTED, align=PP_ALIGN.CENTER)
    add_text_box(slide, left, url_top + Inches(0.5), url_w, Inches(0.4),
                 url, font_size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.4),
             "Built with FastAPI · React · TypeScript · SQLite",
             font_size=12, italic=True, color=MUTED, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
downloads_dir = Path(os.path.expanduser("~")) / "Downloads"
downloads_dir.mkdir(parents=True, exist_ok=True)
output_path = downloads_dir / "WMS_Project_Presentation.pptx"
prs.save(str(output_path))

print(f"OK Saved: {output_path}")
print(f"Size: {output_path.stat().st_size:,} bytes")
print(f"Slides: {TOTAL_SLIDES}")
